"""Office files must never make the converter fetch a URL or read a file named inside the document."""
from __future__ import annotations

import io
import threading
import time
import zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import pytest

from app.converters.sanitize import clean_odf_xml, clean_rels, sanitize_office
from app.errors import ConversionError
from conftest import needs_soffice, post

NS_R = "http://schemas.openxmlformats.org/package/2006/relationships"
IMG = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
LINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"


class _Recorder(BaseHTTPRequestHandler):
    hits: list[str] = []

    def do_GET(self):  # noqa: N802
        type(self).hits.append(self.path)
        self.send_response(404)
        self.end_headers()

    do_HEAD = do_GET  # noqa: N815

    def log_message(self, *a):
        pass


@pytest.fixture()
def listener():
    _Recorder.hits = []
    srv = HTTPServer(("127.0.0.1", 0), _Recorder)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_port}", _Recorder.hits
    srv.shutdown()


def _rewrite(data: bytes, fn) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zout.writestr(info, fn(info.filename, zin.read(info.filename)))
    return buf.getvalue()


def _docx(url: str, suffix: str = "/docx.png") -> bytes:
    from docx import Document

    d = Document()
    d.add_paragraph("hello docx")
    b = io.BytesIO()
    d.save(b)
    pic = (
        '<w:p><w:r><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><wp:extent cx="990000" cy="990000"/><wp:docPr id="1" name="x"/>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:nvPicPr><pic:cNvPr id="0" name="x"/><pic:cNvPicPr/>'
        '</pic:nvPicPr><pic:blipFill><a:blip r:link="rIdExt1"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/>'
        '<a:ext cx="990000" cy="990000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>'
    )

    def fn(name, data):
        if name == "word/document.xml":
            return data.decode().replace("</w:body>", pic + "</w:body>", 1).encode()
        if name == "word/_rels/document.xml.rels":
            return data.decode().replace("</Relationships>", f'<Relationship Id="rIdExt1" Type="{IMG}" Target="{url}{suffix}" TargetMode="External"/></Relationships>').encode()
        return data

    return _rewrite(b.getvalue(), fn)


def _pptx(url: str) -> bytes:
    from pptx import Presentation

    p = Presentation()
    p.slides.add_slide(p.slide_layouts[6])
    b = io.BytesIO()
    p.save(b)
    pic = ('<p:pic><p:nvPicPr><p:cNvPr id="9" name="x"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:link="rIdExt1"/><a:stretch><a:fillRect/></a:stretch>'
           '</p:blipFill><p:spPr><a:xfrm><a:off x="100000" y="100000"/><a:ext cx="2000000" cy="2000000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')

    def fn(name, data):
        if name == "ppt/slides/slide1.xml":
            return data.decode().replace("</p:spTree>", pic + "</p:spTree>", 1).encode()
        if name == "ppt/slides/_rels/slide1.xml.rels":
            return data.decode().replace("</Relationships>", f'<Relationship Id="rIdExt1" Type="{IMG}" Target="{url}/pptx.png" TargetMode="External"/></Relationships>').encode()
        return data

    return _rewrite(b.getvalue(), fn)


def _odf(kind: str, mime: str, body: str) -> bytes:
    content = (
        '<?xml version="1.0" encoding="UTF-8"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:xlink="http://www.w3.org/1999/xlink" '
        'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" office:version="1.2">' + body + "</office:document-content>"
    )
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("mimetype", mime, compress_type=zipfile.ZIP_STORED)
        z.writestr("content.xml", content, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr(
            "META-INF/manifest.xml",
            '<?xml version="1.0"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
            f'<manifest:file-entry manifest:full-path="/" manifest:media-type="{mime}"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>',
        )
    return b.getvalue()


def _odt(url: str) -> bytes:
    return _odf("odt", "application/vnd.oasis.opendocument.text",
                f'<office:body><office:text><text:p>hello odt<draw:frame draw:name="i" text:anchor-type="as-char" svg:width="3cm" svg:height="3cm">'
                f'<draw:image xlink:href="{url}/odt.png" xlink:type="simple" xlink:show="embed" xlink:actuate="onLoad"/></draw:frame></text:p></office:text></office:body>')


def _pdf_text(data: bytes) -> str:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join(p.extract_text() or "" for p in pdf.pages)


# ------------------------------------------------------------------ end to end with a real listener
@needs_soffice
@pytest.mark.parametrize("name,build,marker", [("a.docx", _docx, "hello docx"), ("a.pptx", _pptx, None), ("a.odt", _odt, "hello odt")])
def test_linked_images_are_never_fetched(client, listener, name, build, marker):
    url, hits = listener
    r = post(client, name, build(url), "office-to-pdf", "pdf")
    assert r.status_code == 201, r.text
    time.sleep(0.5)
    assert hits == [], f"LibreOffice fetched {hits}"
    out = client.get(r.json()["downloadUrl"]).content
    assert out[:5] == b"%PDF-"
    if marker:
        assert marker in _pdf_text(out)


@needs_soffice
def test_local_image_files_are_never_read_into_the_pdf(client, tmp_path):
    from PIL import Image

    secret = tmp_path / "secret.png"  # a real, readable image on the server's disk
    Image.new("RGB", (200, 100), "blue").save(secret)
    r = post(client, "a.docx", _docx(f"file://{secret}", ""), "office-to-pdf", "pdf")
    assert r.status_code == 201
    from pypdf import PdfReader

    assert len(PdfReader(io.BytesIO(client.get(r.json()["downloadUrl"]).content)).pages[0].images) == 0


@needs_soffice
def test_embedded_images_and_hyperlinks_still_work(client):
    from docx import Document
    from docx.shared import Inches
    from PIL import Image

    img = io.BytesIO()
    Image.new("RGB", (120, 80), "red").save(img, "PNG")
    img.seek(0)
    d = Document()
    d.add_paragraph("with picture")
    d.add_picture(img, width=Inches(1.5))
    out = io.BytesIO()
    d.save(out)
    r = post(client, "pic.docx", out.getvalue(), "office-to-pdf", "pdf")
    from pypdf import PdfReader

    assert len(PdfReader(io.BytesIO(client.get(r.json()["downloadUrl"]).content)).pages[0].images) == 1


# ------------------------------------------------------------------ sanitiser unit tests
def _rels(*entries: str) -> bytes:
    return f'<?xml version="1.0"?><Relationships xmlns="{NS_R}">{"".join(entries)}</Relationships>'.encode()


def test_rels_keep_internal_parts_and_web_hyperlinks_only():
    data = _rels(
        f'<Relationship Id="1" Type="{IMG}" Target="media/image1.png"/>',
        f'<Relationship Id="2" Type="{IMG}" Target="http://169.254.169.254/latest/meta-data" TargetMode="External"/>',
        f'<Relationship Id="3" Type="{IMG}" Target="file:///etc/passwd" TargetMode="External"/>',
        f'<Relationship Id="4" Type="{IMG}" Target="\\\\attacker\\share\\x.png" TargetMode="External"/>',
        f'<Relationship Id="5" Type="{IMG}" Target="http://x/y.png"/>',  # scheme without TargetMode is still a URL
        f'<Relationship Id="6" Type="{LINK}" Target="https://example.com/page" TargetMode="External"/>',
        f'<Relationship Id="7" Type="{LINK}" Target="file:///etc/passwd" TargetMode="External"/>',
        f'<Relationship Id="8" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate" Target="http://x/t.dotm" TargetMode="External"/>',
        f'<Relationship Id="9" Type="{IMG}" Target="//host/share/x.png"/>',
    )
    out, removed = clean_rels(data, "x.rels")
    assert removed == 7
    text = out.decode()
    assert 'Id="1"' in text and 'Id="6"' in text
    for gone in ("169.254", "file:", "attacker", "http://x/", "attachedTemplate", "//host"):
        assert gone not in text, gone


def test_odf_keeps_internal_paths_and_web_links_and_drops_everything_else():
    xml = (
        '<o xmlns:xlink="http://www.w3.org/1999/xlink" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">'
        '<draw:image xlink:href="Pictures/1.png"/><draw:image xlink:href="http://10.0.0.1/x.png"/><draw:image xlink:href="/etc/passwd"/>'
        '<draw:image xlink:href="../../etc/passwd"/><draw:image xlink:href="file:///etc/passwd"/><table:table-source xlink:href="http://x/s.ods"/>'
        '<text:a xlink:href="https://example.com">link</text:a><text:a xlink:href="vnd.sun.star.script:macro">m</text:a><text:a xlink:href="#bookmark">b</text:a>'
        "</o>"
    ).encode()
    out, removed = clean_odf_xml(xml, "content.xml")
    text = out.decode()
    assert removed == 6
    assert "Pictures/1.png" in text and "https://example.com" in text and "#bookmark" in text
    for gone in ("10.0.0.1", "/etc/passwd", "file:", "x/s.ods", "vnd.sun.star"):
        assert gone not in text, gone


@pytest.mark.parametrize("bad", [
    b'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY u "http://x">]><Relationships xmlns="%s"/>' % NS_R.encode(),
    b"<Relationships",
])
def test_dtds_and_malformed_xml_are_refused(bad):
    with pytest.raises(ConversionError) as e:
        clean_rels(bad, "x.rels")
    assert e.value.code == "invalid_file"


def test_sanitise_enforces_zip_limits_and_rejects_encrypted_or_traversal_entries(tmp_path):
    def make(entries):
        p = tmp_path / "x.docx"
        with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
            for n, d in entries:
                z.writestr(n, d)
        return p

    big = make([("[Content_Types].xml", b"<a/>"), ("word/document.xml", b"0" * 3_000_000)])
    with pytest.raises(ConversionError) as e:
        sanitize_office(big, "docx", max_entries=100, max_expanded_bytes=1_000_000)
    assert e.value.code == "archive_too_large"
    many = make([(f"f{i}.xml", b"<a/>") for i in range(50)])
    with pytest.raises(ConversionError):
        sanitize_office(many, "docx", max_entries=10, max_expanded_bytes=10_000_000)
    trav = make([("../evil.xml", b"<a/>")])
    with pytest.raises(ConversionError):
        sanitize_office(trav, "docx", max_entries=10, max_expanded_bytes=10_000_000)
    notzip = tmp_path / "n.docx"
    notzip.write_bytes(b"PK\x03\x04 not really")
    with pytest.raises(ConversionError):
        sanitize_office(notzip, "docx", max_entries=10, max_expanded_bytes=10_000_000)


def test_odf_mimetype_stays_first_and_stored_after_rewrite(tmp_path):
    p = tmp_path / "x.odt"
    p.write_bytes(_odt("http://127.0.0.1:1"))
    assert sanitize_office(p, "odt", max_entries=100, max_expanded_bytes=10_000_000) == 1
    with zipfile.ZipFile(p) as z:
        first = z.infolist()[0]
        assert first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED
        assert b"127.0.0.1" not in z.read("content.xml")
