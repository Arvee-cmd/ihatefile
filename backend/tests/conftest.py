from __future__ import annotations

import io
import shutil
import subprocess
import tarfile
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from PIL import Image

from app.config import Settings
from app.main import create_app

MARKER = "Hello ihatefile"
HAVE_SOFFICE = bool(shutil.which("soffice") or shutil.which("libreoffice"))
HAVE_PANDOC = bool(shutil.which("pandoc"))
HAVE_FFMPEG = bool(shutil.which("ffmpeg") and shutil.which("ffprobe"))

needs_soffice = pytest.mark.skipif(not HAVE_SOFFICE, reason="LibreOffice not installed")
needs_pandoc = pytest.mark.skipif(not HAVE_PANDOC, reason="Pandoc not installed")
needs_ffmpeg = pytest.mark.skipif(not HAVE_FFMPEG, reason="FFmpeg not installed")


@pytest.fixture()
def settings(tmp_path) -> Settings:
    return replace(Settings(), tmp_root=tmp_path / "tmp", rate_limit_per_minute=100_000, convert_timeout_seconds=120)


@pytest.fixture()
def make_client(settings):
    clients: list[TestClient] = []

    def _make(**overrides) -> TestClient:
        c = TestClient(create_app(replace(settings, **overrides)))
        c.__enter__()
        clients.append(c)
        return c

    yield _make
    for c in clients:
        c.__exit__(None, None, None)


@pytest.fixture()
def client(make_client) -> TestClient:
    return make_client()


def post(client: TestClient, filename: str, data: bytes, tool: str, target: str, options: dict | None = None):
    import json

    return client.post(
        "/api/v1/convert",
        files={"file": (filename, data, "application/octet-stream")},
        data={"tool": tool, "target": target, "options": json.dumps(options or {})},
    )


# ------------------------------------------------------------------ fixture builders
def _png_bytes(size=(64, 48), mode="RGB", color=(200, 30, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new(mode, size, color).save(buf, "PNG")
    return buf.getvalue()


def _run(cmd: list[str]) -> None:
    subprocess.run(cmd, check=True, capture_output=True, timeout=180)


def _ffmpeg(out: Path, *args: str) -> None:
    _run(["ffmpeg", "-hide_banner", "-loglevel", "error", "-y", *args, str(out)])


def _soffice(src: Path, fmt: str) -> Path:
    outdir = src.parent / f"lo_{fmt}"
    outdir.mkdir(exist_ok=True)
    _run(["soffice", f"-env:UserInstallation=file://{src.parent}/lo-prof", "--headless", "--convert-to", fmt, "--outdir", str(outdir), str(src)])
    return outdir / f"{src.stem}.{fmt}"


@pytest.fixture(scope="session")
def samples(tmp_path_factory) -> dict[str, Path]:
    """One small, valid sample file for every input format in the registry."""
    d = tmp_path_factory.mktemp("samples")
    out: dict[str, Path] = {}
    img = Image.new("RGBA", (64, 48), (200, 30, 30, 255))
    for fmt, ext, pil in [("png", "png", "PNG"), ("jpeg", "jpg", "JPEG"), ("webp", "webp", "WEBP"), ("gif", "gif", "GIF"),
                          ("bmp", "bmp", "BMP"), ("tiff", "tiff", "TIFF"), ("ico", "ico", "ICO")]:
        p = d / f"sample.{ext}"
        (img.convert("RGB") if pil in ("JPEG", "BMP") else img).save(p, pil)
        out[fmt] = p
    from reportlab.pdfgen import canvas

    p = d / "sample.pdf"
    c = canvas.Canvas(str(p))
    c.drawString(72, 720, MARKER)
    c.showPage()
    c.drawString(72, 720, "Second page")
    c.save()
    out["pdf"] = p
    out["md"] = d / "sample.md"
    out["md"].write_text(f"# Title\n\n{MARKER} **bold** text.\n\n- one\n- two\n", encoding="utf-8")
    out["html"] = d / "sample.html"
    out["html"].write_text(f"<html><body><h1>Title</h1><p>{MARKER} <b>bold</b> text.</p></body></html>", encoding="utf-8")
    out["rtf"] = d / "sample.rtf"
    out["rtf"].write_text("{\\rtf1\\ansi\\deff0 {\\fonttbl{\\f0 Arial;}}\\f0\\fs24 " + MARKER + " from RTF.\\par}", encoding="ascii")
    out["txt"] = d / "sample.txt"
    out["txt"].write_text(MARKER + "\n", encoding="utf-8")
    if HAVE_PANDOC:
        import docx  # python-docx

        document = docx.Document()
        document.add_heading("Title", 1)
        document.add_paragraph(f"{MARKER} from DOCX.")
        out["docx"] = d / "sample.docx"
        document.save(out["docx"])
        _run(["pandoc", "-f", "markdown", "-t", "odt", "-o", str(d / "sample.odt"), str(out["md"])])
        out["odt"] = d / "sample.odt"
    if HAVE_SOFFICE:
        import openpyxl
        import pptx

        wb = openpyxl.Workbook()
        wb.active["A1"] = MARKER
        out["xlsx"] = d / "sample.xlsx"
        wb.save(out["xlsx"])
        out["ods"] = _soffice(out["xlsx"], "ods")
        deck = pptx.Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = MARKER
        out["pptx"] = d / "sample.pptx"
        deck.save(out["pptx"])
        out["odp"] = _soffice(out["pptx"], "odp")
    if HAVE_FFMPEG:
        sine = ["-f", "lavfi", "-i", "sine=frequency=440:duration=1"]
        _ffmpeg(d / "sample.wav", *sine)
        _ffmpeg(d / "sample.mp3", *sine, "-c:a", "libmp3lame")
        _ffmpeg(d / "sample.flac", *sine)
        _ffmpeg(d / "sample.ogg", *sine, "-c:a", "libvorbis")
        _ffmpeg(d / "sample.m4a", *sine, "-c:a", "aac")
        video = ["-f", "lavfi", "-i", "testsrc=size=160x120:rate=10:duration=1", *sine, "-shortest"]
        _ffmpeg(d / "sample.mp4", *video, "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac")
        _ffmpeg(d / "sample.webm", *video, "-c:v", "libvpx", "-c:a", "libvorbis")
        _ffmpeg(d / "sample.mkv", *video, "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac")
        _ffmpeg(d / "sample.mov", *video, "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-f", "mov")
        for fmt in ("wav", "mp3", "flac", "ogg", "m4a", "mp4", "webm", "mkv", "mov"):
            out[fmt] = d / f"sample.{fmt}"
    with zipfile.ZipFile(d / "sample.zip", "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("a.txt", "alpha")
        zf.writestr("dir/b.txt", "bravo")
    out["zip"] = d / "sample.zip"
    with tarfile.open(d / "sample.tar.gz", "w:gz") as tf:
        for name, body in (("a.txt", b"alpha"), ("dir/b.txt", b"bravo")):
            ti = tarfile.TarInfo(name)
            ti.size = len(body)
            tf.addfile(ti, io.BytesIO(body))
    out["targz"] = d / "sample.tar.gz"
    return out
